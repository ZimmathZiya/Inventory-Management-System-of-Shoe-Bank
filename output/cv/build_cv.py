from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader

OUT = Path(__file__).parent
OUT.mkdir(parents=True, exist_ok=True)
name = 'Mohamed Hajith'
contact = 'Batticaloa, Sri Lanka | +94 74 052 3954 | ahammedlebbehajith@gmail.com'
links = 'github.com/HajithMohamed | linkedin.com/in/mohamed-hajith-b53559295'
summary = 'Full-stack developer and BICT (Hons) undergraduate with project experience in React, Node.js, Express.js, MongoDB, PHP and MySQL. Built web authentication flows, e-commerce features and responsive interfaces, with additional JavaFX team project experience. Seeking a junior full-stack developer or software engineering internship role.'
skills = [
('Languages', 'JavaScript, Java, PHP, SQL, HTML5, CSS3'),
('Frontend', 'React, Redux Toolkit, Tailwind CSS, Bootstrap, jQuery'),
('Backend & databases', 'Node.js, Express.js, Spring Boot, MongoDB, MySQL'),
('Tools & concepts', 'Git, GitHub, Vite, Maven, JavaFX, REST APIs, JWT, OOP')]
projects = [
('Saga Elite | Full-Stack Developer & Team Lead', 'MERN Stack | saga-elite.shop | Team project', [
'Led a team of four to build and deploy an e-commerce platform, from requirements analysis to release.',
'Developed customer and administrator features for product management, JWT authentication and order handling; coordinated GitHub branching workflows.']),
('Faculty Event Management System | Co-Contributor', 'Spring Boot, React, MySQL, JWT, Cloudinary | Team project', [
'Contributed to an event platform with administrator, lecturer and student roles and event approval workflows.',
'Implemented JWT-protected REST APIs, scheduled email reminders and image uploads; collaborated through pull requests and code review.']),
('NEXTGEN | Mobile Shop E-Commerce Website', 'PHP, MySQL, JavaScript, jQuery, Bootstrap', [
'Developed a mobile shop website with product browsing, user accounts, shopping cart and checkout workflows.',
'Built administrative pages for managing products, orders and customers, backed by a MySQL database.']),
('Nano-Zillas | University Management System (Team Project)', 'Java, JavaFX, FXML, MySQL, Maven', [
'Contributed to a JavaFX application with role-specific interfaces for university staff and students.',
'Applied OOP and MVC to modules for courses, attendance, grades and timetables.']),
]

# Compact reference preset with named ATS resume overrides: A4, 0.65-inch
# margins, Arial 10.5pt, black text, compact headings, no page furniture.
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.2677); sec.page_height=Inches(11.6929)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.65)
sec.header_distance=sec.footer_distance=Inches(.3)
def style(n,size,bold=False,before=0,after=3):
 s=doc.styles[n]; s.font.name='Arial'; s.font.size=Pt(size); s.font.bold=bold; s.font.color.rgb=RGBColor(0,0,0)
 p=s.paragraph_format; p.space_before=Pt(before); p.space_after=Pt(after); p.line_spacing=1.08
 return s
style('Normal',10.5);style('Title',23,True,0,3);style('Subtitle',11,True,0,5)
style('Heading 1',11,True,10,5);style('Heading 2',10.5,True,6,2);style('Heading 3',10.5,True,0,2)
style('List Bullet',10.5,False,0,3)
for abstract in doc.part.numbering_part.element.findall(qn('w:abstractNum')):
 for lvl in abstract.findall(qn('w:lvl')):
  fmt=lvl.find(qn('w:numFmt'))
  if fmt is not None and fmt.get(qn('w:val'))=='bullet':
   pr=lvl.find(qn('w:pPr'))
   if pr is None: pr=OxmlElement('w:pPr');lvl.append(pr)
   ind=pr.find(qn('w:ind'))
   if ind is None: ind=OxmlElement('w:ind');pr.append(ind)
   ind.set(qn('w:left'),'220');ind.set(qn('w:hanging'),'220')
doc.add_paragraph(name.upper(),'Title')
doc.add_paragraph('FULL-STACK DEVELOPER','Subtitle')
doc.add_paragraph(contact)
doc.add_paragraph(links)
doc.add_paragraph('PROFESSIONAL SUMMARY','Heading 1');doc.add_paragraph(summary)
doc.add_paragraph('TECHNICAL SKILLS','Heading 1')
for label,text in skills:
 p=doc.add_paragraph();p.add_run(label+': ').bold=True;p.add_run(text)
doc.add_paragraph('PROJECT EXPERIENCE','Heading 1')
for title,stack,bullets in projects:
 doc.add_paragraph(title,'Heading 2')
 p=doc.add_paragraph(stack);p.paragraph_format.space_after=Pt(3)
 for b in bullets:doc.add_paragraph(b,'List Bullet')
doc.add_paragraph('EDUCATION','Heading 1')
doc.add_paragraph('Bachelor of Information and Communication Technology (Hons)','Heading 2')
doc.add_paragraph('University of Ruhuna | Expected completion: 2027')
doc.add_paragraph('Current status: Third year, second semester')
doc.add_paragraph('CERTIFICATION & LANGUAGES','Heading 1')
doc.add_paragraph('The Complete Full-Stack Web Development Bootcamp | Udemy')
doc.add_paragraph('Languages: English, Tamil, Sinhala')
doc.core_properties.author=name;doc.core_properties.title='Mohamed Hajith - Full-Stack Developer CV'
doc.save(OUT/'Mohamed_Hajith_CV.docx')

pdfmetrics.registerFont(TTFont('Arial','C:/Windows/Fonts/arial.ttf'))
pdfmetrics.registerFont(TTFont('Arial-Bold','C:/Windows/Fonts/arialbd.ttf'))
pdfmetrics.registerFontFamily('Arial',normal='Arial',bold='Arial-Bold')
styles={
'name':ParagraphStyle('name',fontName='Arial-Bold',fontSize=23,leading=27,spaceAfter=3),
'role':ParagraphStyle('role',fontName='Arial-Bold',fontSize=11,leading=14,spaceAfter=5),
'body':ParagraphStyle('body',fontName='Arial',fontSize=10.5,leading=13.1,spaceAfter=3),
'contact':ParagraphStyle('contact',fontName='Arial',fontSize=9.5,leading=12.5,spaceAfter=3),
'section':ParagraphStyle('section',fontName='Arial-Bold',fontSize=11,leading=13,spaceBefore=8,spaceAfter=4,keepWithNext=True),
'project':ParagraphStyle('project',fontName='Arial-Bold',fontSize=10.5,leading=13,spaceBefore=6,spaceAfter=2,keepWithNext=True),
'stack':ParagraphStyle('stack',fontName='Arial',fontSize=10,leading=12.5,spaceAfter=3,keepWithNext=True),
'bullet':ParagraphStyle('bullet',fontName='Arial',fontSize=10.5,leading=12.5,leftIndent=11,firstLineIndent=0,bulletIndent=0,spaceAfter=2)}
story=[]
def add(text,kind='body',bullet=None):story.append(Paragraph(text,styles[kind],bulletText=bullet))
add(name.upper(),'name');add('FULL-STACK DEVELOPER','role');add(contact,'contact')
add('<link href="https://github.com/HajithMohamed">github.com/HajithMohamed</link> | <link href="https://www.linkedin.com/in/mohamed-hajith-b53559295">linkedin.com/in/mohamed-hajith-b53559295</link>','contact')
add('PROFESSIONAL SUMMARY','section');add(summary)
add('TECHNICAL SKILLS','section')
for label,text in skills:add('<b>'+escape(label)+':</b> '+escape(text))
add('PROJECT EXPERIENCE','section')
for title,stack,bullets in projects:
 add(title,'project');add(stack,'stack')
 for b in bullets:add(b,'bullet',bullet='\u2022')
add('EDUCATION','section')
add('Bachelor of Information and Communication Technology (Hons)','project')
add('University of Ruhuna | Expected completion: 2027')
add('Current status: Third year, second semester')
add('CERTIFICATION & LANGUAGES','section')
add('The Complete Full-Stack Web Development Bootcamp | Udemy')
add('Languages: English, Tamil, Sinhala')
pdf=OUT/'Mohamed_Hajith_CV.pdf'
SimpleDocTemplate(str(pdf),pagesize=(595.28,841.89),rightMargin=46.8,leftMargin=46.8,topMargin=40,bottomMargin=40,title='Mohamed Hajith - Full-Stack Developer CV',author=name).build(story)
reader=PdfReader(pdf)
assert len(reader.pages)==1, f'Expected 1 page, got {len(reader.pages)}'
txt=reader.pages[0].extract_text()
assert 'ahammedlebbehajith@gmail.com' in txt and 'EDUCATION' in txt
(OUT/'extracted_text.txt').write_text(txt,encoding='utf-8')
print('PDF validated: 1 page; contact details and section text extract correctly.')
